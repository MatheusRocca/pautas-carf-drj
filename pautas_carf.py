# -*- coding: utf-8 -*-
"""
pautas_carf.py
Extrai dados de pautas de julgamento do CARF e DRJ a partir dos XMLs
baixados do INLabs (DOU), filtrando apenas artigos com artType="Pauta".

Aceita uma pasta com XMLs/ZIPs (processa recursivamente) ou um XML direto.
Se rodado sem argumentos, pede os caminhos pelo terminal.

Uso:
    python pautas_carf.py <pasta_ou_xml> <saida.xlsx> [--clientes lista.xlsx]
    python pautas_carf.py ./2026-09-08-DO1 saida.xlsx
    python pautas_carf.py          (menu: baixar o DOU de hoje, outra data ou escolher pasta)
    python pautas_carf.py --baixar             (baixa o DOU de hoje do INLabs e processa)
    python pautas_carf.py --baixar 2026-09-25  (idem, para outra data)
    python pautas_carf.py --baixar --manter-dou  (guarda os ZIPs em DOU/<data>/)
    python pautas_carf.py --instalar-agendamento (roda sozinho em dias úteis, às 8h30, 13h e 18h)
    python pautas_carf.py --testar-aviso         (testa notificação e e-mail)
"""

import sys
import os
import re
import io
import csv
import glob
import argparse
import zipfile
import tempfile
import shutil
import subprocess
import unicodedata
import time
import json
import plistlib
import smtplib
import traceback
import configparser
import http.client
import http.cookiejar
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from html import escape
from html.parser import HTMLParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')


# ─────────────────────────────────────────────
# Extração de texto HTML
# ─────────────────────────────────────────────

class TextExtractor(HTMLParser):
    """
    Junta o texto do HTML ignorando os parágrafos de assinatura
    (<p class="assina"> / <p class="cargo">): sem isso, o nome do presidente
    da turma gruda no último processo da pauta quando vem em maiúsculas.
    """
    CLASSES_IGNORADAS = {'assina', 'cargo'}

    def __init__(self):
        super().__init__()
        self.parts = []
        self.ignorando = False
    def handle_starttag(self, tag, attrs):
        if tag == 'p':
            classes = (dict(attrs).get('class') or '').split()
            self.ignorando = any(c in self.CLASSES_IGNORADAS for c in classes)
    def handle_endtag(self, tag):
        if tag == 'p':
            self.ignorando = False
    def handle_data(self, data):
        if not self.ignorando:
            self.parts.append(data)
    def get_text(self):
        return ' '.join(self.parts)

def html_para_texto(html):
    p = TextExtractor()
    p.feed(html or '')
    return re.sub(r'\s+', ' ', p.get_text()).strip()


# ─────────────────────────────────────────────
# Descompactação de ZIPs
# ─────────────────────────────────────────────

def extrair_zips(pasta, destino):
    """Extrai ZIPs encontrados na pasta para destino (recursivo, aninhado)."""
    processados = set()
    pendentes = glob.glob(os.path.join(pasta, '**', '*.zip'), recursive=True)
    contador = 0

    while pendentes:
        zpath = pendentes.pop()
        real = os.path.abspath(zpath)
        if real in processados:
            continue
        processados.add(real)
        if len(processados) > 5000:
            print("  [AVISO] Limite de 5000 ZIPs atingido — parando extração.")
            break

        contador += 1
        subdestino = os.path.join(destino, f"_zip{contador:04d}")
        try:
            with zipfile.ZipFile(zpath, 'r') as zf:
                zf.extractall(subdestino)
            print(f"  [ZIP] Extraído: {os.path.basename(zpath)}")
        except zipfile.BadZipFile:
            print(f"  [ERRO] ZIP inválido: {os.path.basename(zpath)}")
            continue

        novos = glob.glob(os.path.join(subdestino, '**', '*.zip'), recursive=True)
        pendentes.extend(n for n in novos if os.path.abspath(n) not in processados)

    return contador


# ─────────────────────────────────────────────
# Parser de XML de pauta
# ─────────────────────────────────────────────

# Marca o início de cada processo numerado e captura o número do processo.
PAT_PROCESSO_MARCADOR = re.compile(
    r'\d+\s*-\s*Processo\s+n[oº°]?[:\s]+([0-9./\-]+)\s*-\s*',
    re.IGNORECASE
)

# Rótulo ("Recorrente", "Embargante"...) seguido de valor, até o próximo
# rótulo ou até o texto virar nome próprio em caixa mista (assinatura do
# presidente da turma, que vem sem rótulo ao final da pauta).
PAT_LABEL_VALOR = re.compile(
    r'([A-Za-zÀ-ÿ()]{3,40}?)\s*:\s*(.+?)'
    r'(?=\s+(?:e\s+)?[A-Za-zÀ-ÿ()]{3,40}?\s*:'
    r'|\s+(?:e\s+)?[A-ZÀ-Ý][a-zà-ÿ]'
    r'|\Z)'
)

ROLES_RECORRENTE = {
    'recorrente', 'recorrentes', 'embargante', 'embargantes', 'requerente',
    'requerentes', 'impugnante', 'impugnantes', 'reclamante', 'reclamantes',
    'autuado', 'autuada', 'autuados', 'autuadas', 'autor', 'autora', 'autores',
}
ROLES_INTERESSADO = {
    'interessado', 'interessada', 'interessados', 'interessadas',
    'embargado', 'embargada', 'embargados', 'embargadas',
    'recorrido', 'recorrida', 'recorridos', 'recorridas',
    'reclamado', 'reclamada', 'reclamados', 'reclamadas',
    'requerido', 'requerida', 'requeridos', 'requeridas',
}

PAT_DIA = re.compile(
    r'\bDIA\s+(\d{1,2}\s+de\s+[A-Za-zÀ-ÿ]+\s+de\s+\d{4}),?\s*[ÀA]S\s+(\d{2}:\d{2})\s*HORAS?',
    re.IGNORECASE
)


def extrair_processos(texto_proc):
    """
    Divide um bloco de texto em processos numerados ('N - Processo nº: ...')
    e extrai Recorrente/Interessado tolerando rótulos variados (Recorrentes,
    Embargante, Recorrido, Requerente...) e mais de duas partes por processo.
    Qualquer rótulo não reconhecido cai em Recorrente/Interessado (o que
    estiver vazio) ou, se os dois já tiverem valor, em 'outras_partes'.
    """
    marcadores = list(PAT_PROCESSO_MARCADOR.finditer(texto_proc))
    resultados = []
    for i, m in enumerate(marcadores):
        inicio = m.end()
        fim = marcadores[i + 1].start() if i + 1 < len(marcadores) else len(texto_proc)
        bloco = texto_proc[inicio:fim].strip()

        recorrente_partes, interessado_partes, outras_partes = [], [], []
        for lm in PAT_LABEL_VALOR.finditer(bloco):
            label = re.sub(r'\(.*?\)', '', lm.group(1)).strip().lower()
            valor = lm.group(2).strip().rstrip('.,;')
            if not valor:
                continue
            if label in ROLES_RECORRENTE:
                recorrente_partes.append(valor)
            elif label in ROLES_INTERESSADO:
                interessado_partes.append(valor)
            elif not recorrente_partes:
                recorrente_partes.append(valor)
            elif not interessado_partes:
                interessado_partes.append(valor)
            else:
                outras_partes.append(f"{label.capitalize()}: {valor}")

        resultados.append({
            'numero':        m.group(1).strip(),
            'recorrente':    ' / '.join(recorrente_partes),
            'interessado':   ' / '.join(interessado_partes),
            'outras_partes': ' / '.join(outras_partes),
        })
    return resultados


# Órgãos cujas pautas o script deve conseguir ler (CARF e DRJ). Uma pauta
# deles sem nenhum processo extraído indica provável mudança de formato.
ORGAOS_MONITORADOS = ('Conselho Administrativo de Recursos Fiscais', 'Receita Federal')


PAT_PARTE_XML = re.compile(r'^(?P<base>.+)-(?P<parte>\d+)\.xml$', re.IGNORECASE)


def agrupar_partes(xmls):
    """
    Pautas grandes vêm do INLabs divididas em vários XMLs, um por página
    ('..._24451873-1.xml', '-2', ...). Só a 1ª parte (e as que abrem uma
    nova sessão) tem o cabeçalho com data, horário e relator; as outras
    continuam a lista. Agrupa as partes de cada ato, em ordem, para serem
    lidas como um texto só. XMLs sem sufixo de parte ficam sozinhos.
    """
    grupos = {}
    for xml in xmls:
        m = PAT_PARTE_XML.match(os.path.basename(xml))
        chave = (os.path.dirname(xml), m.group('base')) if m else (xml, None)
        grupos.setdefault(chave, []).append((int(m.group('parte')) if m else 0, xml))
    return [[x for _, x in sorted(partes)] for _, partes in sorted(grupos.items())]


def parse_xml_pauta(filepath, avisos=None):
    """
    Lê um XML do INLabs — ou a lista de partes de um mesmo ato, em ordem
    (ver agrupar_partes) — e retorna lista de dicts (um por processo).
    Retorna [] se não for pauta de julgamento.
    Problemas encontrados são impressos e, se `avisos` for uma lista,
    também acrescentados a ela (para o resumo no final da execução).
    """
    def avisar(msg, rotulo='[AVISO]'):
        print(f"  {rotulo} {msg}")
        if avisos is not None:
            avisos.append(msg)

    partes_xml = [filepath] if isinstance(filepath, str) else list(filepath)
    artigos = []
    for parte in partes_xml:
        try:
            art = ET.parse(parte).getroot().find('article')
        except ET.ParseError as e:
            avisar(f"XML inválido: {os.path.basename(parte)}: {e}", '[ERRO]')
            return []
        if art is None:
            return []
        artigos.append(art)
    filepath = partes_xml[0]
    art = artigos[0]

    # Filtra só artigos do tipo Pauta
    art_type = art.get('artType', '')
    if 'pauta' not in art_type.lower():
        return []

    # Metadados
    categoria = art.get('artCategory', '')
    pub_date  = art.get('pubDate', '')
    pdf_page  = art.get('pdfPage', '')
    nome_arq  = os.path.basename(filepath)

    partes = [p.strip() for p in categoria.split('/')]
    orgao  = ' / '.join(partes[-2:]) if len(partes) >= 2 else categoria
    pauta  = partes[-1] if partes else ''

    texto = ' '.join(html_para_texto(a.findtext('.//Texto') or '') for a in artigos)
    if len(partes_xml) > 1:
        nome_arq = f"{nome_arq} + {len(partes_xml) - 1} parte(s)"
    esperados = len(PAT_PROCESSO_MARCADOR.findall(texto))

    linhas = []
    blocos_dia = list(PAT_DIA.finditer(texto))

    # texto antes do 1º cabeçalho de sessão (ou o texto todo, se não houver):
    # processos sem data/horário — não podem ser descartados
    antes = texto[:blocos_dia[0].start()] if blocos_dia else texto
    linhas = _extrair_do_bloco(antes, orgao, pauta, '', '', pub_date, pdf_page or nome_arq)
    if blocos_dia:
        for idx, m_dia in enumerate(blocos_dia):
            data_julg = m_dia.group(1).strip().rstrip(',')
            horario   = m_dia.group(2).strip()
            inicio    = m_dia.end()
            fim       = blocos_dia[idx + 1].start() if idx + 1 < len(blocos_dia) else len(texto)
            bloco     = texto[inicio:fim]
            linhas.extend(_extrair_do_bloco(bloco, orgao, pauta, data_julg, horario, pub_date, pdf_page or nome_arq))

    onde = f"{orgao} ({nome_arq}{' — ' + pdf_page if pdf_page else ''})"
    if len(linhas) != esperados:
        avisar(f"{esperados} processo(s) numerado(s) no texto, mas {len(linhas)} extraído(s) — "
               f"possível perda de dados. Confira no DOU: {onde}")
    elif not linhas and any(o in categoria for o in ORGAOS_MONITORADOS):
        avisar(f"pauta do CARF/DRJ sem nenhum processo reconhecido — o formato pode ter mudado. "
               f"Confira no DOU: {onde}")

    return linhas


def _extrair_do_bloco(bloco, orgao, pauta, data_julg, horario, pub_date, arquivo):
    """
    Divide um bloco (um DIA de sessão, ou o texto todo quando não há
    cabeçalho de DIA) pelos rótulos 'Relator(a):' — que no formato do CARF
    precedem o(s) processo(s) a que se referem — e extrai os processos de
    cada sub-bloco com o relator correspondente.
    """
    linhas = []
    sub_blocos = re.split(r'Relator\(a\):\s*', bloco, flags=re.IGNORECASE)
    sem_relator_declarado = len(sub_blocos) == 1
    relator_atual = ''

    for sub in sub_blocos:
        sub = sub.strip()
        if not sub:
            continue
        m_primeiro = re.search(r'\d+\s*-\s*Processo', sub, re.IGNORECASE)
        if not m_primeiro:
            relator_atual = sub.strip()
            continue
        if not sem_relator_declarado:
            relator_atual = sub[:m_primeiro.start()].strip().rstrip(',.:')
        texto_proc = sub[m_primeiro.start():]

        for proc in extrair_processos(texto_proc):
            linhas.append(_montar_linha(orgao, pauta, proc, data_julg, horario, relator_atual, pub_date, arquivo))
    return linhas


def _montar_linha(orgao, pauta, proc, data_julg, horario, relator, pub_date, arquivo):
    return {
        'Órgão':               orgao,
        'Pauta':               pauta,
        'Número do Processo':  proc['numero'],
        'Recorrente':          proc['recorrente'],
        'Interessado':         proc['interessado'],
        'Outras Partes':       proc['outras_partes'],
        'Data de Julgamento':  data_julg,
        'Horário':             horario,
        'Relator(a)':          relator,
        'Data Publicação DOU': pub_date,
        'Link DOU':            arquivo,
    }


# ─────────────────────────────────────────────
# Checagem de clientes na pauta
# ─────────────────────────────────────────────
#
# Duas checagens independentes, combinadas para dar mais confiança:
#
# 1) Número do processo: se a planilha de controle já lista o número do
#    processo, é uma chave exata (sem ambiguidade) — a comparação mais
#    confiável possível.
# 2) Nome da parte: as pautas do CARF/DRJ nunca trazem CNPJ/CPF, só nome —
#    então esse é um match por aproximação de texto, menos confiável.
#
# Quando as duas concordam, o resultado é marcado como "confirmado". Um
# processo cadastrado na planilha mas sem nome vinculado (comum: apensados
# de controle interno) ainda conta como match — só não tem como confirmar
# pelo nome.

COLUNAS_NOME_CLIENTE     = {'nome', 'cliente', 'razao social', 'nome cliente', 'nome do cliente', 'empresa',
                             'partes e empresas monitoradas', 'partes empresas monitoradas'}
COLUNAS_DOC_CLIENTE      = {'cnpj', 'cpf', 'cnpj cpf', 'cpf cnpj', 'documento'}
COLUNAS_PROCESSO_CLIENTE = {'processo', 'numero do processo', 'n do processo', 'processos',
                             'processo controle interno', 'processo controle', 'no processo'}

SUFIXOS_SOCIETARIOS = re.compile(r'\b(LTDA|LTD|EIRELI|ME|EPP|SA|S A)\b')


def normalizar_nome(nome):
    """Maiúsculas, sem acento, só letras/números/espaço."""
    n = unicodedata.normalize('NFKD', str(nome or ''))
    n = ''.join(c for c in n if not unicodedata.combining(c))
    n = n.upper()
    n = re.sub(r'[^A-Z0-9 ]', ' ', n)
    return re.sub(r'\s+', ' ', n).strip()


def nome_base(nome_normalizado):
    """Remove sufixos societários (LTDA, S/A, ME...) de um nome já normalizado."""
    n = SUFIXOS_SOCIETARIOS.sub('', nome_normalizado)
    return re.sub(r'\s+', ' ', n).strip()


def _norm_cabecalho(s):
    n = unicodedata.normalize('NFKD', str(s or ''))
    n = ''.join(c for c in n if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', ' ', n.lower()).strip()


def _achar_coluna(cabecalho, candidatas):
    for i, h in enumerate(cabecalho):
        if _norm_cabecalho(h) in candidatas:
            return i
    return None


def _desmesclar(ws):
    """
    Retorna a matriz de valores da planilha (linhas x colunas), propagando
    o valor da célula-mãe para todas as células de cada intervalo mesclado.
    Sem isso, openpyxl devolve None para as células "escondidas" atrás de
    uma mesclagem — comum em listas de controle onde um número de processo
    mesclado vale para várias linhas/empresas.
    """
    valores = [[c.value for c in row] for row in ws.iter_rows()]
    for rng in ws.merged_cells.ranges:
        v = valores[rng.min_row - 1][rng.min_col - 1]
        for r in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                valores[r - 1][col - 1] = v
    return valores


def carregar_clientes(caminho):
    """
    Lê uma lista de clientes/processos (.xlsx/.csv) e retorna:
      - clientes: lista de dicts (nome, doc, nome_base) para o match por nome
      - processos_index: dict {número de processo (só dígitos) -> lista de
        nomes vinculados a ele na planilha, podendo ser vazia quando o
        processo está cadastrado mas sem nome de cliente (ex.: apensados de
        controle interno)}

    Detecta as colunas de nome, documento e número de processo por
    cabeçalhos comuns — não exige nomes exatos de coluna.
    """
    ext = os.path.splitext(caminho)[1].lower()

    if ext in ('.xlsx', '.xlsm'):
        wb = load_workbook(caminho, data_only=True)
        valores = _desmesclar(wb.active)
        cabecalho = valores[0] if valores else None
        linhas_raw = valores[1:]
    elif ext == '.csv':
        with open(caminho, newline='', encoding='utf-8-sig') as fp:
            todas = list(csv.reader(fp))
        cabecalho = todas[0] if todas else None
        linhas_raw = todas[1:]
    else:
        raise ValueError(f"Formato não suportado para lista de clientes: '{ext}' (use .xlsx ou .csv)")

    if not cabecalho:
        raise ValueError(f"Lista de clientes vazia: {caminho}")

    idx_nome = _achar_coluna(cabecalho, COLUNAS_NOME_CLIENTE)
    idx_doc  = _achar_coluna(cabecalho, COLUNAS_DOC_CLIENTE)
    idx_proc = _achar_coluna(cabecalho, COLUNAS_PROCESSO_CLIENTE)
    if idx_nome is None:
        raise ValueError(
            f"Não encontrei coluna de nome do cliente em {caminho} "
            f"(cabeçalho lido: {list(cabecalho)}). Renomeie a coluna para "
            "'Nome' ou 'Razão Social'."
        )

    clientes = []
    processos_index = {}
    for row in linhas_raw:
        nome = str(row[idx_nome] or '').strip() if idx_nome < len(row) else ''
        if nome:
            clientes.append({
                'nome': nome,
                'doc': str(row[idx_doc]).strip() if idx_doc is not None and idx_doc < len(row) and row[idx_doc] else '',
                'nome_base': nome_base(normalizar_nome(nome)),
            })

        if idx_proc is not None and idx_proc < len(row) and row[idx_proc]:
            proc_digitos = re.sub(r'\D', '', str(row[idx_proc]))
            if proc_digitos:
                nomes = processos_index.setdefault(proc_digitos, [])
                if nome and nome not in nomes:
                    nomes.append(nome)

    return clientes, processos_index


def encontrar_cliente_por_nome(recorrente, interessado, outras_partes, clientes):
    """
    Procura, entre os clientes cadastrados, algum cujo nome apareça no
    Recorrente/Interessado/Outras Partes do processo. Compara por nome
    normalizado (sem acento/maiúsculas/sufixo societário) com fronteira de
    palavra, nas duas direções — tanto a pauta quanto a lista de clientes
    podem trazer o nome completo ou abreviado. Esse é o match "de reserva":
    menos confiável que o número do processo, por ser aproximação de texto.
    """
    texto_partes = ' '.join(p for p in (recorrente, interessado, outras_partes) if p)
    campo_base = nome_base(normalizar_nome(texto_partes))
    if not campo_base:
        return None, None

    for cliente in clientes:
        base = cliente['nome_base']
        if len(base) < 4:
            continue  # nome curto demais para comparar com segurança
        if (base == campo_base
                or re.search(r'\b' + re.escape(base) + r'\b', campo_base)
                or (len(campo_base) >= 6 and re.search(r'\b' + re.escape(campo_base) + r'\b', base))):
            return cliente['nome'], _trecho_que_bateu(base, (recorrente, interessado, outras_partes), texto_partes)
    return None, None


def _trecho_que_bateu(base, campos, texto_partes):
    """Qual parte da pauta (ex.: um dos recorrentes) contém o nome do cliente."""
    for campo in campos:
        for parte in (campo or '').split(' / '):
            parte_base = nome_base(normalizar_nome(parte))
            if parte_base and (re.search(r'\b' + re.escape(base) + r'\b', parte_base)
                               or re.search(r'\b' + re.escape(parte_base) + r'\b', base)):
                return parte.strip()
    return texto_partes


def checar_cliente(numero_processo, recorrente, interessado, outras_partes, clientes, processos_index):
    """
    Faz as duas checagens (número do processo e nome da parte) e combina o
    resultado em (nome_identificado, tipo_de_match, motivo). Prioriza o número do
    processo — é uma chave exata — e usa o nome como confirmação extra ou,
    quando o processo não está cadastrado, como único critério.
    """
    numero_digitos = re.sub(r'\D', '', numero_processo or '')
    nomes_processo = processos_index.get(numero_digitos) if numero_digitos else None
    nome_por_texto, trecho = encontrar_cliente_por_nome(recorrente, interessado, outras_partes, clientes)
    partes_pauta = ' / '.join(p for p in (recorrente, interessado) if p)

    if nomes_processo is not None:
        if nomes_processo:
            if nome_por_texto in nomes_processo:
                return (' / '.join(nomes_processo), 'Número do processo + nome (confirmado)',
                        f'Nº do processo está na lista; na pauta "{trecho}" = na lista "{nome_por_texto}"')
            return (' / '.join(nomes_processo), 'Número do processo (confira o nome)',
                    f'Nº do processo está na lista como "{" / ".join(nomes_processo)}", '
                    f'mas na pauta a parte é "{partes_pauta}"')
        # processo cadastrado no controle, mas sem nome de cliente vinculado
        return (partes_pauta, 'Número do processo (sem nome cadastrado no controle)',
                'Nº do processo está na lista, sem nome cadastrado')

    if nome_por_texto:
        return (nome_por_texto, 'Nome (sem número cadastrado no controle — confira)',
                f'Só pelo nome: na pauta "{trecho}" ↔ na lista "{nome_por_texto}"')

    return '', None, ''


# ─────────────────────────────────────────────
# Geração do Excel
# ─────────────────────────────────────────────

COLUNAS_BASE = [
    'Órgão', 'Pauta', 'Número do Processo', 'Recorrente',
    'Interessado', 'Outras Partes', 'Data de Julgamento', 'Horário',
    'Relator(a)', 'Data Publicação DOU', 'Link DOU',
]
COLUNAS_CLIENTE = ['Cliente Identificado', 'Tipo de Match', 'Por que bateu']

COLS_ALINHADAS_ESQUERDA = {
    'Órgão', 'Pauta', 'Recorrente', 'Interessado', 'Outras Partes',
    'Relator(a)', 'Cliente Identificado', 'Tipo de Match', 'Por que bateu',
}

LARGURAS = {
    'Órgão': 35, 'Pauta': 30, 'Número do Processo': 26,
    'Recorrente': 45, 'Interessado': 22, 'Outras Partes': 30,
    'Data de Julgamento': 14, 'Horário': 9,
    'Relator(a)': 32, 'Data Publicação DOU': 14, 'Link DOU': 14,
    'Cliente Identificado': 30, 'Tipo de Match': 28, 'Por que bateu': 60,
}

ILLEGAL = re.compile(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]')

COLS_DATA = {'Data de Julgamento', 'Data Publicação DOU'}

MESES = {
    'janeiro': 1, 'fevereiro': 2, 'marco': 3, 'abril': 4, 'maio': 5, 'junho': 6,
    'julho': 7, 'agosto': 8, 'setembro': 9, 'outubro': 10, 'novembro': 11, 'dezembro': 12,
}


def texto_para_data(texto):
    """
    Converte '7 de Outubro de 2026' (data da sessão) ou '25/09/2026' (data
    do DOU) para date, para o Excel ordenar e filtrar como data. Devolve o
    texto original se não reconhecer o formato — melhor mostrar o texto do
    que perder a informação.
    """
    t = (texto or '').strip()
    m = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', t)
    if m:
        dia, mes, ano = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = re.fullmatch(r'(\d{1,2})\s+de\s+(\S+)\s+de\s+(\d{4})', t, re.IGNORECASE)
        mes = MESES.get(normalizar_nome(m.group(2)).lower()) if m else None
        if not mes:
            return texto
        dia, ano = int(m.group(1)), int(m.group(3))
    try:
        return date(ano, mes, dia)
    except ValueError:
        return texto


THIN     = Side(style='thin', color='BFBFBF')
BORDA    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_HDR = PatternFill('solid', start_color='D9E1F2')
# cor da linha conforme a confiança do match com a lista de clientes
FILL_CONFIRMADO = PatternFill('solid', start_color='63BE7B')  # verde forte: número + nome
FILL_NUMERO     = PatternFill('solid', start_color='C6EFCE')  # verde claro: só o número
FILL_TALVEZ     = PatternFill('solid', start_color='FFEB9C')  # amarelo: só o nome
FONT_HDR  = Font(name='Calibri', bold=True, size=9)
FONT_N    = Font(name='Calibri', size=9)
FONT_LINK = Font(name='Calibri', size=9, color='0563C1', underline='single')
CENTER    = Alignment(horizontal='center', vertical='center')
LEFT      = Alignment(horizontal='left',   vertical='center')


def _preencher_sheet(ws, titulo, colunas, linhas, marcar_cliente):
    if titulo:
        ws.title = titulo

    for j, col in enumerate(colunas, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.font = FONT_HDR; c.fill = FILL_HDR; c.border = BORDA; c.alignment = CENTER
    ws.row_dimensions[1].height = 28

    for i, linha in enumerate(linhas, start=2):
        fill_linha = None
        if marcar_cliente and linha.get('Cliente Identificado'):
            tipo = linha.get('Tipo de Match', '')
            if 'confirmado' in tipo:
                fill_linha = FILL_CONFIRMADO
            elif tipo.startswith('Número do processo'):
                fill_linha = FILL_NUMERO
            else:
                fill_linha = FILL_TALVEZ
        for j, col in enumerate(colunas, start=1):
            valor = linha.get(col, '')
            if isinstance(valor, str):
                valor = ILLEGAL.sub('', valor)
            link = None
            if col in COLS_DATA:
                valor = texto_para_data(valor)
            elif col == 'Link DOU' and valor.startswith('http'):
                link, valor = valor, 'Abrir no DOU'
            c = ws.cell(row=i, column=j, value=valor)
            c.font = FONT_N; c.border = BORDA
            c.alignment = LEFT if col in COLS_ALINHADAS_ESQUERDA else CENTER
            if isinstance(valor, date):
                c.number_format = 'DD/MM/YYYY'
            if link:
                c.hyperlink = link
                c.font = FONT_LINK
            if fill_linha:
                c.fill = fill_linha

    for j, col in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(j)].width = LARGURAS.get(col, 15)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{len(linhas) + 1}"


def _preencher_legenda(ws):
    """Aba separada com a legenda das cores (fora da tabela, para não atrapalhar o filtro)."""
    legenda = [
        ('Legenda das cores', None, None),
        ('Número do processo + nome (confirmado)', FILL_CONFIRMADO,
         'O nº do processo está na lista de clientes e o nome da parte também bate.'),
        ('Só o número do processo', FILL_NUMERO,
         'O nº do processo está na lista, mas o nome na pauta é diferente ou não há nome cadastrado.'),
        ('Só o nome — conferir', FILL_TALVEZ,
         'O nº do processo não está na lista; bateu só pela semelhança do nome. Confira.'),
    ]
    for i, (texto, fill, explicacao) in enumerate(legenda, start=1):
        c = ws.cell(row=i, column=1, value=texto)
        c.alignment = LEFT
        if fill is None:
            c.font = FONT_HDR
            continue
        c.font = FONT_N; c.fill = fill; c.border = BORDA
        e = ws.cell(row=i, column=2, value=explicacao)
        e.font = FONT_N; e.alignment = LEFT
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 90


def gerar_xlsx(todas_linhas, caminho_saida, com_clientes=False):
    wb = Workbook()
    colunas = COLUNAS_BASE + (COLUNAS_CLIENTE if com_clientes else [])

    _preencher_sheet(wb.active, "Pautas", colunas, todas_linhas, marcar_cliente=com_clientes)

    if com_clientes:
        linhas_cliente = [l for l in todas_linhas if l.get('Cliente Identificado')]
        if linhas_cliente:
            ws2 = wb.create_sheet('Clientes na Pauta')
            _preencher_sheet(ws2, None, colunas, linhas_cliente, marcar_cliente=True)
            print(f"   {len(linhas_cliente)} processo(s) com cliente identificado — ver aba 'Clientes na Pauta'.")
        _preencher_legenda(wb.create_sheet('Legenda'))

    wb.save(caminho_saida)
    print(f"\n[OK] Planilha salva em: {caminho_saida}")
    print(f"   {len(todas_linhas)} processos extraídos.")


# ─────────────────────────────────────────────
# Download automático do INLabs
# ─────────────────────────────────────────────
#
# Segue o script oficial da Imprensa Nacional
# (github.com/Imprensa-Nacional/inlabs): login com e-mail/senha, que devolve
# o cookie 'inlabs_session_cookie', e um ZIP de XMLs por seção do DOU.
# Baixa todas as seções — as pautas costumam sair no DO1, mas as outras
# vêm junto para não arriscar perder uma pauta publicada fora do lugar.

URL_LOGIN    = 'https://inlabs.in.gov.br/logar.php'
URL_DOWNLOAD = 'https://inlabs.in.gov.br/index.php?p={data}&dl={data}-{secao}.zip'
SECOES_DOU     = ['DO1', 'DO1E', 'DO2', 'DO2E', 'DO3', 'DO3E']
SECOES_EXTRAS  = {'DO1E', 'DO2E', 'DO3E'}
TENTATIVAS     = 3
TENTATIVAS_LOGIN = 5  # o login do INLabs às vezes só entra depois de várias tentativas
ESPERA_ENTRE_TENTATIVAS = 10  # segundos (multiplicado pelo nº da tentativa)
ARQ_CREDENCIAIS = 'inlabs.ini'

MODELO_CREDENCIAIS = """\
# Login do INLabs (https://inlabs.in.gov.br) usado para baixar o DOU.
# Este arquivo fica só neste computador — não vai para o GitHub.
[inlabs]
email =
senha =
"""


class ErroInlabs(Exception):
    pass


def ler_credenciais(pasta):
    """
    E-mail e senha do INLabs: das variáveis INLABS_EMAIL/INLABS_SENHA, ou do
    arquivo inlabs.ini na pasta do script (criado vazio na primeira vez).
    """
    email, senha = os.environ.get('INLABS_EMAIL'), os.environ.get('INLABS_SENHA')
    if email and senha:
        return email, senha

    caminho = os.path.join(pasta, ARQ_CREDENCIAIS)
    if not os.path.exists(caminho):
        with open(caminho, 'w', encoding='utf-8') as fp:
            fp.write(MODELO_CREDENCIAIS)
        os.chmod(caminho, 0o600)
        raise ErroInlabs(f"Criei o arquivo {caminho} — preencha o e-mail e a senha do INLabs nele e rode de novo.")

    cp = configparser.ConfigParser(interpolation=None)  # sem interpolação: senha pode ter '%'
    cp.read(caminho, encoding='utf-8')
    email = cp.get('inlabs', 'email', fallback='').strip()
    senha = cp.get('inlabs', 'senha', fallback='').strip()
    if not email or not senha:
        raise ErroInlabs(f"Preencha o e-mail e a senha do INLabs em {caminho} e rode de novo.")
    return email, senha


# --- Lista de clientes no Google Sheets (usado ao rodar no GitHub Actions) ---
# Com CLIENTES_PLANILHA_ID definido, a lista é baixada do Google Sheets em vez
# de procurada na pasta. O acesso é por uma conta de serviço do Google (JSON em
# GOOGLE_CREDENCIAIS) com quem a planilha foi compartilhada como leitor.

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def baixar_clientes_google(planilha_id, destino):
    """Exporta a planilha como .xlsx (mantém as células mescladas, que
    carregar_clientes já sabe tratar) e devolve o caminho salvo."""
    from google.oauth2 import service_account          # só instalado no GitHub Actions
    from google.auth.transport.requests import Request
    try:
        info = json.loads(os.environ.get('GOOGLE_CREDENCIAIS', ''))
    except ValueError:
        raise ValueError("GOOGLE_CREDENCIAIS vazio ou não é o JSON da conta de serviço")
    cred = service_account.Credentials.from_service_account_info(
        info, scopes=['https://www.googleapis.com/auth/drive.readonly'])
    cred.refresh(Request())
    url = (f'https://www.googleapis.com/drive/v3/files/{urllib.parse.quote(planilha_id)}/export'
           f'?mimeType={urllib.parse.quote(XLSX_MIME)}')
    req = urllib.request.Request(url, headers={'Authorization': f'Bearer {cred.token}'})
    try:
        with urllib.request.urlopen(req, timeout=60) as resp, open(destino, 'wb') as fp:
            shutil.copyfileobj(resp, fp)
    except urllib.error.HTTPError as e:
        if e.code in (403, 404):
            raise ValueError(f"sem acesso à planilha de clientes do Google (HTTP {e.code}) — confira o ID e se "
                             f"ela foi compartilhada com {info.get('client_email')}") from e
        raise
    return destino


def _requisitar(opener, req, tentativas=TENTATIVAS):
    """
    Faz a requisição, repetindo em falha de rede ou erro do servidor (5xx).
    Devolve (status, conteúdo); 4xx volta direto, sem repetir.
    """
    ultimo_erro = None
    for tentativa in range(1, tentativas + 1):
        try:
            with opener.open(req, timeout=120) as r:
                return r.status, r.read()
        except urllib.error.HTTPError as e:
            if e.code < 500:
                return e.code, b''
            ultimo_erro = f'servidor respondeu HTTP {e.code}'
        except (urllib.error.URLError, http.client.HTTPException, OSError) as e:
            ultimo_erro = str(getattr(e, 'reason', None) or e) or type(e).__name__
        if tentativa < tentativas:
            time.sleep(ESPERA_ENTRE_TENTATIVAS * tentativa)
    raise ErroInlabs(f'falhou após {tentativas} tentativas ({ultimo_erro})')


class _SemRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):
        return None


def login_inlabs(email, senha):
    jar = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    # O login responde com redirecionamento (302) e já traz o cookie da
    # sessão; não segue o redirecionamento — a página seguinte não é
    # necessária e é mais um ponto que pode falhar.
    opener_login = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar), _SemRedirect())
    req = urllib.request.Request(
        URL_LOGIN,
        data=urllib.parse.urlencode({'email': email, 'password': senha}).encode(),
        headers={'Content-Type': 'application/x-www-form-urlencoded'},
    )
    # o INLabs às vezes falha ou responde sem a sessão mesmo com a senha
    # certa, e entra numa tentativa seguinte: repete o login inteiro
    ultimo_erro = None
    for tentativa in range(1, TENTATIVAS_LOGIN + 1):
        try:
            _requisitar(opener_login, req, tentativas=1)
            if any(c.name == 'inlabs_session_cookie' for c in jar):
                return opener
            ultimo_erro = None
        except ErroInlabs as e:
            ultimo_erro = e
        if tentativa < TENTATIVAS_LOGIN:
            time.sleep(ESPERA_ENTRE_TENTATIVAS * tentativa)
    if ultimo_erro:
        raise ErroInlabs(f'Não consegui acessar o INLabs para fazer login '
                         f'(falhou após {TENTATIVAS_LOGIN} tentativas): {ultimo_erro}')
    raise ErroInlabs(f'O INLabs recusou o login {TENTATIVAS_LOGIN} vezes — '
                     'confira o e-mail e a senha em inlabs.ini.')


def _tamanho(caminho):
    n = os.path.getsize(caminho)
    if n < 1_000_000:
        return f"{max(1, round(n / 1000))} KB"
    return f"{n / 1_000_000:.1f} MB".replace('.', ',')


def baixar_dou(data_iso, destino, pasta_credenciais, secoes=SECOES_DOU):
    """
    Baixa os ZIPs do DOU de uma data para `destino`. Devolve uma lista de
    (seção, situação, detalhe), com situação 'OK', 'SEM_EXTRA' (edição extra
    que não saiu — normal), 'NAO_PUBLICADO' (seção principal sem arquivo)
    ou 'ERRO'. Um ZIP de seção principal já baixado e íntegro não é baixado
    de novo; as edições extras sempre são, porque podem sair mais ao longo
    do dia (e são pequenas).
    Falha de login levanta ErroInlabs (aí nada foi baixado).
    """
    os.makedirs(destino, exist_ok=True)
    opener = login_inlabs(*ler_credenciais(pasta_credenciais))

    resultados = []
    for secao in secoes:
        caminho = os.path.join(destino, f'{data_iso}-{secao}.zip')
        if secao not in SECOES_EXTRAS and zipfile.is_zipfile(caminho):
            resultados.append((secao, 'OK', f'já estava baixado ({_tamanho(caminho)})'))
            continue

        print(f"  Baixando {secao}...")
        req = urllib.request.Request(URL_DOWNLOAD.format(data=data_iso, secao=secao),
                                     headers={'origem': '736372697074'})  # cabeçalho do script oficial
        try:
            status, conteudo = _requisitar(opener, req)
        except ErroInlabs as e:
            resultados.append((secao, 'ERRO', str(e)))
            continue

        if status == 404 or _pagina_sem_a_data(conteudo, data_iso):
            if secao in SECOES_EXTRAS:
                resultados.append((secao, 'SEM_EXTRA', 'sem edição extra nesse dia'))
            else:
                resultados.append((secao, 'NAO_PUBLICADO',
                                   'não publicado (fim de semana, feriado ou ainda não saiu)'))
            continue
        if status != 200:
            resultados.append((secao, 'ERRO', f'o INLabs respondeu HTTP {status}'))
            continue

        parcial = caminho + '.part'
        with open(parcial, 'wb') as fp:
            fp.write(conteudo)
        if not zipfile.is_zipfile(parcial):
            os.remove(parcial)
            resultados.append((secao, 'ERRO', 'o INLabs não devolveu um ZIP válido (arquivo corrompido ou sessão expirada)'))
            continue
        os.replace(parcial, caminho)
        resultados.append((secao, 'OK', _tamanho(caminho)))

    return resultados


def _pagina_sem_a_data(conteudo, data_iso):
    """Para uma data sem DOU (fim de semana, feriado) o INLabs não dá 404:
    devolve a página logada com a lista de datas disponíveis. Distingue isso
    da página de login (sessão expirada, que continua sendo erro)."""
    if zipfile.is_zipfile(io.BytesIO(conteudo)) or not conteudo.lstrip()[:15].lower().startswith((b'<!doctype', b'<html')):
        return False
    pagina = conteudo.lower()
    logado = b'type="password"' not in pagina and b"type='password'" not in pagina and b'sair' in pagina
    return logado and data_iso.encode() not in conteudo


ROTULOS_SITUACAO = {'OK': '[OK]   ', 'SEM_EXTRA': '[--]   ', 'NAO_PUBLICADO': '[AVISO]', 'ERRO': '[ERRO] '}


def imprimir_resumo_download(data_iso, resultados):
    print(f"\nDownload do DOU de {_data_br(data_iso)} (INLabs):")
    for secao, situacao, detalhe in resultados:
        print(f"  {ROTULOS_SITUACAO[situacao]} {secao:<5} {detalhe}")


def _apagar_download(pasta, manter):
    """Apaga os ZIPs baixados de um dia (e a pasta DOU/, se ficar vazia)."""
    if manter:
        if os.path.isdir(pasta):
            print(f"\n  ZIPs do DOU mantidos em: {pasta}")
        return
    shutil.rmtree(pasta, ignore_errors=True)
    try:
        os.rmdir(os.path.dirname(pasta))  # só apaga DOU/ se estiver vazia
    except OSError:
        pass
    print("\n  ZIPs do DOU apagados (use --manter-dou para guardar).")


def _data_br(data_iso):
    return datetime.strptime(data_iso, '%Y-%m-%d').strftime('%d/%m/%Y')


def interpretar_data(texto):
    """Aceita 'AAAA-MM-DD' ou 'DD/MM/AAAA'; devolve 'AAAA-MM-DD'."""
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(texto.strip(), fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    raise argparse.ArgumentTypeError(f"data inválida: '{texto}' (use AAAA-MM-DD ou DD/MM/AAAA)")


# ─────────────────────────────────────────────
# Seletores (janela nativa do macOS, com fallback no terminal)
# ─────────────────────────────────────────────
#
# Não usa Tkinter: no macOS, o Tcl/Tk 8.5 que vem com o Python do sistema
# tem um bug em que a janela de escolha simplesmente não aparece. Em vez
# disso usa o seletor nativo do Finder via osascript. Se não der (outro
# sistema, erro), pede o caminho no terminal — dá pra arrastar a pasta.

def _osascript(script):
    if sys.platform != 'darwin':
        return None
    try:
        r = subprocess.run(['osascript', '-e', script], capture_output=True, text=True)
    except OSError:
        return None
    if r.returncode != 0:
        # código 1 com "User canceled" (-128) = usuário cancelou
        return '' if '-128' in r.stderr else None
    return r.stdout.strip()


def _pedir_caminho(mensagem):
    caminho = input(mensagem).strip()
    # remove aspas que o Terminal adiciona ao arrastar um arquivo/pasta
    if len(caminho) >= 2 and caminho[0] == caminho[-1] and caminho[0] in ('"', "'"):
        caminho = caminho[1:-1]
    # o zsh escapa espaços ao arrastar (ex.: Meu\ Drive)
    caminho = caminho.replace('\\ ', ' ')
    return caminho or None


def selecionar_pasta():
    res = _osascript('POSIX path of (choose folder with prompt '
                     '"Selecione a pasta com os XMLs/ZIPs de pauta")')
    if res is not None:
        return res or None
    return _pedir_caminho('Pasta com os XMLs/ZIPs de pauta (ou caminho de um XML) — '
                           'digite ou arraste pro terminal: ')


def selecionar_saida(nome='pautas.xlsx', pasta=None):
    script = f'POSIX path of (choose file name with prompt "Salvar planilha como" default name "{nome}"'
    if pasta and os.path.isdir(pasta):
        script += f' default location (POSIX file "{pasta}")'
    script += ')'
    res = _osascript(script)
    if res is not None:
        if res and not res.lower().endswith('.xlsx'):
            res += '.xlsx'
        return res or None
    return _pedir_caminho(f"Arquivo Excel de saída [Enter para '{nome}' na mesma pasta]: ")


def imprimir_problemas(erros_download, problema_clientes, avisos_parsing, data_baixada):
    """
    Repete no final da execução, em destaque, tudo o que pode ter deixado a
    planilha incompleta — para não passar despercebido no meio da saída.
    Devolve a lista de problemas (vazia se não houve nenhum).
    """
    problemas = ([f"{secao} não foi baixado: {detalhe}" for secao, detalhe in erros_download]
                 + ([f"clientes NÃO foram checados: {problema_clientes}"] if problema_clientes else [])
                 + [f"problema de leitura: {aviso}" for aviso in avisos_parsing])
    if not problemas:
        return problemas
    print("\n" + "!" * 70)
    print("[ATENÇÃO] A planilha pode estar incompleta:")
    if erros_download:
        print(f"\n  {len(erros_download)} seção(ões) do DOU NÃO foram baixadas — a planilha NÃO inclui as pautas delas:")
        for secao, detalhe in erros_download:
            print(f"    - {secao}: {detalhe}")
        print(f"    Rode de novo para tentar baixar de novo:  "
              f"python3 \"{os.path.basename(__file__)}\" --baixar {data_baixada}")
    if problema_clientes:
        print(f"\n  Clientes NÃO foram checados: {problema_clientes}.")
    if avisos_parsing:
        print(f"\n  {len(avisos_parsing)} pauta(s) com problema de leitura:")
        for aviso in avisos_parsing:
            print(f"    - {aviso}")
    print("!" * 70)
    return problemas


def menu_inicial():
    """
    Menu mostrado quando o script roda sem opções. Devolve a data a baixar
    ('AAAA-MM-DD') ou None para escolher uma pasta já baixada no Finder.
    """
    hoje = date.today()
    print("Pautas CARF/DRJ — o que fazer?")
    print(f"  [Enter]     baixar o DOU de hoje ({hoje.strftime('%d/%m/%Y')})")
    print("  DD/MM/AAAA  baixar o DOU de outra data")
    print("  p           escolher uma pasta com XMLs/ZIPs já baixados")
    while True:
        try:
            resposta = input("> ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\nEncerrando.")
            sys.exit(0)
        if not resposta:
            return hoje.isoformat()
        if resposta.lower() == 'p':
            return None
        try:
            return interpretar_data(resposta)
        except argparse.ArgumentTypeError as e:
            print(f"  {e} — tente de novo.")


# ─────────────────────────────────────────────
# Execução
# ─────────────────────────────────────────────

@dataclass
class Resultado:
    """O que uma execução produziu — usado pelo modo automático para avisar."""
    codigo: int = 0                 # 0 ok, 1 nada a processar, 2 algum problema
    planilha: str = None
    pautas: dict = field(default_factory=dict)   # XML da pauta → nº de processos
    linhas_cliente: list = field(default_factory=list)
    problemas: list = field(default_factory=list)


def executar(args, pasta_script):
    """
    Baixa (se args.baixar), processa os XMLs e gera a planilha. Imprime o
    andamento e devolve um Resultado em vez de encerrar o programa.
    """
    res = Resultado()
    erros_download = []
    if args.baixar:
        saida = args.entrada  # com --baixar, o único argumento posicional é a saída
        args.entrada = os.path.join(pasta_script, 'DOU', args.baixar)
        args.saida = saida or os.path.join(pasta_script, f'pautas_{args.baixar}.xlsx')

        print(f"Baixando o DOU de {_data_br(args.baixar)} do INLabs para {args.entrada}")
        try:
            resultados = baixar_dou(args.baixar, args.entrada, pasta_script)
        except ErroInlabs as e:
            print(f"\n[ERRO] {e}")
            _apagar_download(args.entrada, args.manter_dou)
            res.codigo, res.problemas = 2, [str(e)]
            return res
        imprimir_resumo_download(args.baixar, resultados)
        erros_download = [(secao, detalhe) for secao, situacao, detalhe in resultados if situacao == 'ERRO']
        if not any(situacao == 'OK' for _, situacao, _ in resultados):
            print(f"\nNenhuma seção do DOU de {_data_br(args.baixar)} foi baixada — nada a processar.")
            _apagar_download(args.entrada, args.manter_dou)
            res.problemas = [f"{secao} não foi baixado: {detalhe}" for secao, detalhe in erros_download]
            res.codigo = 2 if erros_download else 1
            return res
        print()

    if not args.entrada:
        print("Abrindo seletor de pasta...")
        args.entrada = selecionar_pasta()
        if not args.entrada:
            print("Nenhuma pasta selecionada. Encerrando.")
            return res

    if not os.path.exists(args.entrada):
        print(f"Caminho não encontrado: {args.entrada}")
        res.codigo = 1
        return res

    if not args.saida:
        print("Abrindo seletor de destino...")
        base = args.entrada if os.path.isdir(args.entrada) else os.path.dirname(args.entrada)
        args.saida = selecionar_saida(pasta=base)
        if not args.saida:
            base = args.entrada if os.path.isdir(args.entrada) else os.path.dirname(args.entrada)
            args.saida = os.path.join(base, 'pautas.xlsx')
            print(f"Salvando em: {args.saida}")

    avisos_parsing = []
    problema_clientes = None

    # Pasta temporária para ZIPs — mantida durante todo o processamento
    tmpdir = tempfile.mkdtemp(prefix='pauta_zip_')
    try:
        if os.path.isfile(args.entrada):
            xmls = [args.entrada]
        else:
            zips = glob.glob(os.path.join(args.entrada, '**', '*.zip'), recursive=True)
            if zips:
                print(f"  Encontrados {len(zips)} ZIP(s) — extraindo...")
                extrair_zips(args.entrada, tmpdir)

            xmls  = glob.glob(os.path.join(args.entrada, '**', '*.xml'), recursive=True)
            xmls += glob.glob(os.path.join(tmpdir,       '**', '*.xml'), recursive=True)
            xmls  = sorted(set(xmls))

        if not xmls:
            print(f"Nenhum XML encontrado em: {args.entrada}")
            res.codigo = 1
            return res

        print(f"Encontrados {len(xmls)} XMLs. Filtrando pautas...")

        todas_linhas = []
        pautas_encontradas = 0
        for partes in agrupar_partes(xmls):
            linhas = parse_xml_pauta(partes, avisos_parsing)
            if linhas:
                pautas_encontradas += 1
                nome = os.path.basename(partes[0])
                extra = f" (+{len(partes) - 1} parte(s))" if len(partes) > 1 else ''
                print(f"  [PAUTA] {nome}{extra} → {len(linhas)} processo(s)")
                res.pautas[nome] = len(linhas)
                todas_linhas.extend(linhas)

        print(f"\n  {pautas_encontradas} pauta(s) encontrada(s) de {len(xmls)} XMLs.")

        if not todas_linhas:
            print("  Nenhuma pauta de julgamento do CARF/DRJ nestes XMLs — a planilha vai sair vazia.")

        clientes, processos_index = [], {}
        planilha_google = os.environ.get('CLIENTES_PLANILHA_ID', '').strip()
        if not args.clientes and planilha_google:
            try:
                args.clientes = baixar_clientes_google(
                    planilha_google, os.path.join(tmpdir, 'Clientes monitorados (Google Sheets).xlsx'))
            except Exception as e:
                problema_clientes = f"não consegui baixar a lista de clientes do Google Sheets ({e})"
        elif not args.clientes:
            # sem --clientes: usa a lista "Clientes monitorados*.xlsx" que
            # estiver na pasta do script (a mais recente, se houver várias)
            candidatas = [c for c in glob.glob(os.path.join(pasta_script, 'Clientes monitorados*.xlsx'))
                          if not os.path.basename(c).startswith('~$')]
            if candidatas:
                args.clientes = max(candidatas, key=os.path.getmtime)
            else:
                problema_clientes = ("nenhuma lista de clientes informada nem encontrada na pasta do script "
                                     "(esperado: 'Clientes monitorados*.xlsx')")
        if args.clientes:
            if not os.path.exists(args.clientes):
                problema_clientes = f"lista de clientes não encontrada: {args.clientes}"
            else:
                try:
                    clientes, processos_index = carregar_clientes(args.clientes)
                    print(f"  {len(clientes)} cliente(s) e {len(processos_index)} processo(s) de controle "
                          f"carregados de {args.clientes}")
                    if not (clientes or processos_index):
                        problema_clientes = f"a lista de clientes está vazia: {args.clientes}"
                except Exception as e:
                    problema_clientes = f"não consegui ler a lista de clientes ({e})"
        if problema_clientes:
            print(f"[AVISO] {problema_clientes} — seguindo SEM checagem de clientes.")

        com_clientes = bool(clientes or processos_index)
        if com_clientes:
            for linha in todas_linhas:
                nome, tipo, motivo = checar_cliente(
                    linha.get('Número do Processo', ''), linha.get('Recorrente', ''),
                    linha.get('Interessado', ''), linha.get('Outras Partes', ''),
                    clientes, processos_index,
                )
                linha['Cliente Identificado'] = nome
                linha['Tipo de Match'] = tipo or ''
                linha['Por que bateu'] = motivo
                if nome:
                    res.linhas_cliente.append(linha)
            print(f"  {len(res.linhas_cliente)} processo(s) com cliente identificado na pauta.")

        gerar_xlsx(todas_linhas, args.saida, com_clientes=com_clientes)
        res.planilha = args.saida

    except Exception:
        # erro inesperado: mantém o download para dar para investigar
        if args.baixar:
            print(f"\n[AVISO] Os ZIPs baixados foram mantidos em {args.entrada} para investigar o erro abaixo.")
            args.manter_dou = True
        raise
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
        if args.baixar:
            _apagar_download(args.entrada, args.manter_dou)
        res.problemas = imprimir_problemas(erros_download, problema_clientes, avisos_parsing, args.baixar)

    if res.problemas:
        res.codigo = 2
    return res


# ─────────────────────────────────────────────
# Modo automático (agendado pelo launchd do macOS)
# ─────────────────────────────────────────────
#
# Roda sozinho em dias úteis, às 8h30, 13h e 18h (se o Mac estiver dormindo
# nesse horário, o macOS roda assim que ele acordar). Cada execução processa hoje
# e os dias anteriores ainda não concluídos (até DIAS_RECUPERAR para trás),
# então um dia em que o Mac ficou desligado é recuperado na próxima vez, e
# o dia de ontem é reprocessado na manhã seguinte para pegar edições
# extras publicadas no fim do dia. A primeira execução do dia sempre manda
# e-mail (com a planilha); as seguintes só mandam se saiu pauta nova ainda
# não vista, se apareceu cliente novo ou se algo deu errado.

ARQ_ESTADO      = 'automatico_estado.json'
ARQ_EMAIL       = 'email.ini'
PASTA_LOGS      = 'logs'
DIAS_RECUPERAR  = 7
DIAS_HISTORICO  = 60   # quanto tempo o estado lembra dos avisos já enviados
ROTULO_LAUNCHD  = 'local.pautas-carf-drj'
HORARIOS        = [(8, 30), (13, 0), (18, 0)]
DIAS_UTEIS      = [1, 2, 3, 4, 5]  # segunda a sexta (convenção do launchd)
# A partir desta hora o e-mail vai só para a conta que envia (não para a
# equipe) e o estado não é salvo: a execução da manhã seguinte refaz tudo e
# avisa todo mundo. Acontece quando o GitHub atrasa a execução das 18h.
HORA_SO_PARA_MIM = 20

MODELO_EMAIL = """\
# Conta usada para ENVIAR os avisos do modo automático.
# Este arquivo fica só neste computador — não vai para o GitHub.
#
# Gmail / Google Workspace: servidor smtp.gmail.com, porta 465, e uma
# "senha de app" (não é a senha normal da conta) — gere em
# https://myaccount.google.com/apppasswords (precisa da verificação em 2 etapas).
[email]
servidor = smtp.gmail.com
porta = 465
usuario =
senha =
# Quem recebe os avisos (separe vários e-mails por vírgula)
destinatarios =
"""


def carregar_estado(pasta, hoje):
    caminho = os.path.join(pasta, ARQ_ESTADO)
    try:
        with open(caminho, encoding='utf-8') as fp:
            estado = json.load(fp)
    except FileNotFoundError:
        # primeira execução: começa hoje, sem recuperar dias antigos
        estado = {'inicio': hoje.isoformat()}
    estado.setdefault('inicio', hoje.isoformat())
    estado.setdefault('concluidos', [])
    estado.setdefault('avisados', {})
    estado.setdefault('pautas_vistas', {})
    estado.setdefault('problemas_avisados', {})
    return estado


def salvar_estado(pasta, estado, hoje):
    limite = (hoje - timedelta(days=DIAS_HISTORICO)).isoformat()
    estado['concluidos'] = sorted(d for d in set(estado['concluidos']) if d >= limite)
    estado['avisados'] = {d: v for d, v in estado['avisados'].items() if d >= limite}
    estado['pautas_vistas'] = {d: v for d, v in estado.get('pautas_vistas', {}).items() if d >= limite}
    # problemas só valem para o dia: amanhã o mesmo problema é avisado de novo
    estado['problemas_avisados'] = {d: v for d, v in estado.get('problemas_avisados', {}).items()
                                    if d == hoje.isoformat()}
    caminho = os.path.join(pasta, ARQ_ESTADO)
    with open(caminho + '.tmp', 'w', encoding='utf-8') as fp:
        json.dump(estado, fp, ensure_ascii=False, indent=1)
    os.replace(caminho + '.tmp', caminho)


def datas_pendentes(estado, hoje):
    """Hoje e os dias anteriores (até DIAS_RECUPERAR) ainda não concluídos."""
    inicio = max(date.fromisoformat(estado['inicio']), hoje - timedelta(days=DIAS_RECUPERAR))
    pendentes, d = [], inicio
    while d <= hoje:
        if d.isoformat() not in estado['concluidos']:
            pendentes.append(d.isoformat())
        d += timedelta(days=1)
    return pendentes


def rodar_automatico(pasta_script, hoje=None, agora=None):
    hoje = hoje or date.today()
    agora = agora or datetime.now()
    noite = agora.hour >= HORA_SO_PARA_MIM
    print("=" * 70)
    print(f"Execução automática — {agora.strftime('%d/%m/%Y %H:%M')}")
    if noite:
        print(f"Depois das {HORA_SO_PARA_MIM}h: e-mail só para a conta que envia, e o estado não é salvo.")
    estado = carregar_estado(pasta_script, hoje)
    pendentes = datas_pendentes(estado, hoje)
    print(f"Dias a processar: {', '.join(_data_br(d) for d in pendentes)}")

    novos_clientes, problemas, planilhas, novas_pautas = [], [], [], []
    for data_iso in pendentes:
        print(f"\n----- {_data_br(data_iso)} -----")
        args = argparse.Namespace(entrada=None, saida=None, clientes=None, baixar=data_iso, manter_dou=False)
        try:
            res = executar(args, pasta_script)
        except Exception as e:
            traceback.print_exc()
            problemas.append((data_iso, f"erro inesperado: {type(e).__name__}: {e}"))
            continue
        problemas.extend((data_iso, p) for p in res.problemas)

        ja_avisados = set(estado['avisados'].get(data_iso, []))
        novos = [l for l in res.linhas_cliente if l['Número do Processo'] not in ja_avisados]
        novos_clientes.extend((data_iso, l) for l in novos)
        estado['avisados'][data_iso] = sorted(ja_avisados | {l['Número do Processo'] for l in res.linhas_cliente})

        vistas = set(estado['pautas_vistas'].get(data_iso, []))
        novas = {x: n for x, n in res.pautas.items() if x not in vistas}
        if novas:
            novas_pautas.append((data_iso, len(novas), sum(novas.values())))
        estado['pautas_vistas'][data_iso] = sorted(vistas | set(res.pautas))

        if (novos or novas) and res.planilha:
            planilhas.append(res.planilha)

        # o dia só é dado como concluído depois que terminou e rodou sem problemas
        if data_iso < hoje.isoformat() and not res.problemas:
            estado['concluidos'].append(data_iso)

    print(f"\nResumo: {len(novos_clientes)} cliente(s) novo(s) na pauta, "
          f"{sum(n for _, n, _ in novas_pautas)} pauta(s) nova(s), {len(problemas)} problema(s).")
    primeiro_do_dia = estado.get('resumo_enviado') != hoje.isoformat()
    # o mesmo problema só é avisado uma vez por dia (o agendamento tenta várias vezes)
    ja_avisados = estado['problemas_avisados'].get(hoje.isoformat(), [])
    problemas_novos = [p for p in problemas if f"{p[0]} {p[1]}" not in ja_avisados]
    # a equipe só recebe aviso com cliente ou pauta nova, ou o resumo do dia se
    # tudo deu certo; aviso só de erro vai só para a conta que envia, e não
    # conta como resumo: ele sai para a equipe quando o problema se resolver
    para_equipe = novos_clientes or novas_pautas or (primeiro_do_dia and not problemas)
    if para_equipe or problemas_novos:
        enviar_avisos(pasta_script, novos_clientes, problemas_novos, novas_pautas, planilhas, hoje, primeiro_do_dia,
                      so_para_mim=noite or not para_equipe)
        if para_equipe:
            estado['resumo_enviado'] = hoje.isoformat()
        estado['problemas_avisados'][hoje.isoformat()] = ja_avisados + [f"{d} {p}" for d, p in problemas_novos]
    else:
        print("Nada novo desde o último aviso.")
    if not noite:
        salvar_estado(pasta_script, estado, hoje)
    return 2 if problemas else 0


def _titulo_aviso(novos_clientes, problemas, novas_pautas, hoje, primeiro_do_dia):
    if not novos_clientes and not problemas:
        if primeiro_do_dia:
            return f"Pautas CARF/DRJ de {_data_br(hoje.isoformat())}: nenhum cliente na pauta"
        n = sum(n for _, n, _ in novas_pautas)
        return f"Pautas CARF/DRJ de {_data_br(hoje.isoformat())}: {n} pauta(s) nova(s), nenhum cliente"
    partes = []
    if novos_clientes:
        partes.append(f"{len(novos_clientes)} cliente(s) na pauta")
    if problemas:
        partes.append(f"{len(problemas)} problema(s)")
    return "Pautas CARF/DRJ: " + " e ".join(partes)


def _texto_aviso(novos_clientes, problemas, novas_pautas, pasta_script, com_anexo):
    linhas = []
    if not novos_clientes and not problemas:
        linhas.append("Nenhum cliente monitorado apareceu nas pautas novas." if novas_pautas
                      else "Não houve pautas do CARF/DRJ no DOU de hoje.")
    if novos_clientes:
        linhas.append("CLIENTES NA PAUTA")
        for data_iso, l in novos_clientes:
            julgamento = texto_para_data(l.get('Data de Julgamento', ''))
            julgamento = julgamento.strftime('%d/%m/%Y') if isinstance(julgamento, date) else julgamento
            linhas += [
                "",
                f"• {l['Cliente Identificado']}",
                f"  Processo: {l['Número do Processo']}",
                f"  Partes na pauta: {l['Recorrente']} × {l['Interessado']}",
                f"  Órgão: {l['Órgão']}",
                f"  Julgamento: {julgamento} {l.get('Horário', '')} — relator(a) {l.get('Relator(a)', '')}",
                f"  Como bateu: {l['Tipo de Match']}",
                f"  DOU de {_data_br(data_iso)}: {l.get('Link DOU', '')}",
            ]
    if problemas:
        if linhas:
            linhas.append("")
        linhas.append("PROBLEMAS (a planilha pode estar incompleta)")
        for data_iso, p in problemas:
            linhas.append(f"• DOU de {_data_br(data_iso)}: {p}")
    if novas_pautas:
        linhas += ["", "PAUTAS NOVAS DESDE O ÚLTIMO AVISO"]
        linhas += [f"• DOU de {_data_br(d)}: {n} pauta(s), {p} processo(s)" for d, n, p in novas_pautas]
    if com_anexo:
        linhas += ["", "Planilha(s) atualizada(s) em anexo."]
    linhas += ["", _onde_ficam(pasta_script)]
    return "\n".join(linhas)


def _onde_ficam(pasta_script):
    if os.environ.get('GITHUB_RUN_ID'):
        return ("Log e planilhas desta execução: "
                f"{os.environ.get('GITHUB_SERVER_URL', 'https://github.com')}/"
                f"{os.environ.get('GITHUB_REPOSITORY', '')}/actions/runs/{os.environ['GITHUB_RUN_ID']}")
    return f"Planilhas e log em: {pasta_script}"


def _html_aviso(titulo, novos_clientes, problemas, novas_pautas, pasta_script, com_anexo):
    """Mesmo conteúdo de _texto_aviso, formatado para o e-mail (estilos inline,
    que é o que os leitores de e-mail respeitam)."""
    e = lambda v: escape(str(v or ''))
    th = 'style="text-align:left;padding:6px 10px;background:#eef1f5;border-bottom:1px solid #d5dae1;font-size:12px;color:#44506a"'
    td = 'style="padding:8px 10px;border-bottom:1px solid #e6e9ee;vertical-align:top;font-size:13px"'
    secao = 'style="margin:22px 0 8px;font-size:14px;color:#1f2a44;text-transform:uppercase;letter-spacing:.04em"'
    partes = [f'<h2 style="margin:0 0 4px;font-size:18px;color:#1f2a44">{e(titulo)}</h2>']

    if not novos_clientes and not problemas:
        msg = ("Nenhum cliente monitorado apareceu nas pautas novas." if novas_pautas
               else "Não houve pautas do CARF/DRJ no DOU de hoje.")
        partes.append(f'<p style="margin:12px 0;padding:10px 12px;background:#eef7ee;border-left:4px solid #3c8d4a">{e(msg)}</p>')

    if novos_clientes:
        partes.append(f'<h3 {secao}>Clientes na pauta</h3>')
        linhas = []
        for data_iso, l in novos_clientes:
            julgamento = texto_para_data(l.get('Data de Julgamento', ''))
            julgamento = julgamento.strftime('%d/%m/%Y') if isinstance(julgamento, date) else julgamento
            link = l.get('Link DOU', '')
            dou = (f'<a href="{e(link)}" style="color:#2456a6">DOU de {_data_br(data_iso)}</a>' if link
                   else f'DOU de {_data_br(data_iso)}')
            linhas.append(
                f'<tr><td {td}><b>{e(l["Cliente Identificado"])}</b><br>'
                f'<span style="color:#5a6478;font-size:12px">{e(l["Recorrente"])} × {e(l["Interessado"])}</span></td>'
                f'<td {td}>{e(l["Número do Processo"])}<br><span style="color:#5a6478;font-size:12px">{e(l["Órgão"])}</span></td>'
                f'<td {td}><b>{e(julgamento)}</b> {e(l.get("Horário", ""))}<br>'
                f'<span style="color:#5a6478;font-size:12px">{e(l.get("Relator(a)", ""))}</span></td>'
                f'<td {td}>{e(l["Tipo de Match"])}<br><span style="font-size:12px">{dou}</span></td></tr>')
        partes.append(f'<table cellspacing="0" cellpadding="0" style="border-collapse:collapse;width:100%">'
                      f'<tr><th {th}>Cliente / partes</th><th {th}>Processo / órgão</th>'
                      f'<th {th}>Julgamento / relator(a)</th><th {th}>Como bateu</th></tr>{"".join(linhas)}</table>')

    if problemas:
        itens = ''.join(f'<li>DOU de {_data_br(d)}: {e(p)}</li>' for d, p in problemas)
        partes.append(f'<h3 {secao}>Problemas</h3>'
                      f'<div style="padding:10px 12px;background:#fdf1ef;border-left:4px solid #c0392b">'
                      f'<b>A planilha pode estar incompleta.</b><ul style="margin:6px 0 0;padding-left:18px">{itens}</ul></div>')

    if novas_pautas:
        linhas = ''.join(f'<tr><td {td}>DOU de {_data_br(d)}</td><td {td}>{n}</td><td {td}>{p}</td></tr>'
                         for d, n, p in novas_pautas)
        partes.append(f'<h3 {secao}>Pautas novas desde o último aviso</h3>'
                      f'<table cellspacing="0" cellpadding="0" style="border-collapse:collapse">'
                      f'<tr><th {th}>Publicação</th><th {th}>Pautas</th><th {th}>Processos</th></tr>{linhas}</table>')

    rodape = (["Planilha(s) atualizada(s) em anexo."] if com_anexo else []) + [_onde_ficam(pasta_script)]
    partes.append('<p style="margin:24px 0 0;color:#7a8396;font-size:12px">' + '<br>'.join(e(r) for r in rodape) + '</p>')
    return ('<!doctype html><html><body style="margin:0;padding:16px;background:#ffffff">'
            '<div style="max-width:760px;font-family:Arial,Helvetica,sans-serif;color:#1f2a44;line-height:1.4">'
            + ''.join(partes) + '</div></body></html>')


def enviar_avisos(pasta_script, novos_clientes, problemas, novas_pautas, planilhas, hoje, primeiro_do_dia,
                  so_para_mim=False):
    titulo = _titulo_aviso(novos_clientes, problemas, novas_pautas, hoje, primeiro_do_dia)
    corpo = _texto_aviso(novos_clientes, problemas, novas_pautas, pasta_script, bool(planilhas))
    html = _html_aviso(titulo, novos_clientes, problemas, novas_pautas, pasta_script, bool(planilhas))
    print(f"\nAvisando: {titulo}")

    erro_email = None
    try:
        config = ler_config_email(pasta_script)
        if so_para_mim:
            config = dict(config, destinatarios=[config['usuario']])
        enviar_email(config, titulo, corpo, planilhas, html=html)
        print(f"  E-mail enviado para {', '.join(config['destinatarios'])}.")
    except Exception as e:
        erro_email = str(e)
        print(f"  [ERRO] E-mail não enviado: {erro_email}")

    resumo = []
    if novos_clientes:
        nomes = sorted({l['Cliente Identificado'] for _, l in novos_clientes})
        resumo.append(', '.join(nomes))
    if problemas:
        resumo.append(problemas[0][1] + (f" (+{len(problemas) - 1})" if len(problemas) > 1 else ''))
    if not resumo and not erro_email:
        resumo.append("Resumo enviado por e-mail.")
    if erro_email:
        resumo.append(f"E-mail NÃO enviado: {erro_email}")
    notificar_mac(titulo, ' | '.join(resumo))


def notificar_mac(titulo, mensagem):
    if sys.platform != 'darwin':
        return
    # texto passado como argumento (não embutido no AppleScript), para
    # aspas e acentos no nome dos clientes não quebrarem o comando
    script = ['on run argv',
              'display notification (item 2 of argv) with title (item 1 of argv) sound name "Glass"',
              'end run']
    cmd = ['osascript'] + [a for linha in script for a in ('-e', linha)] + [titulo, mensagem[:250]]
    try:
        subprocess.run(cmd, capture_output=True, timeout=30)
        print("  Notificação mostrada no Mac.")
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"  [ERRO] Notificação não mostrada: {e}")


def ler_config_email(pasta):
    if os.environ.get('EMAIL_USUARIO'):
        # no GitHub Actions a conta vem dos segredos, não do email.ini
        return {
            'servidor': os.environ.get('EMAIL_SERVIDOR') or 'smtp.gmail.com',
            'porta': int(os.environ.get('EMAIL_PORTA') or 465),
            'usuario': os.environ['EMAIL_USUARIO'].strip(),
            'senha': os.environ.get('EMAIL_SENHA', '').strip(),
            'destinatarios': [d.strip() for d in os.environ.get('EMAIL_DESTINATARIOS', '').split(',') if d.strip()],
        }
    caminho = os.path.join(pasta, ARQ_EMAIL)
    if not os.path.exists(caminho):
        with open(caminho, 'w', encoding='utf-8') as fp:
            fp.write(MODELO_EMAIL)
        os.chmod(caminho, 0o600)
        raise ValueError(f"criei {caminho} — preencha a conta de envio e os destinatários")
    cp = configparser.ConfigParser(interpolation=None)
    cp.read(caminho, encoding='utf-8')
    config = {
        'servidor': cp.get('email', 'servidor', fallback='').strip(),
        'porta': cp.getint('email', 'porta', fallback=465),
        'usuario': cp.get('email', 'usuario', fallback='').strip(),
        'senha': cp.get('email', 'senha', fallback='').strip(),
        'destinatarios': [d.strip() for d in cp.get('email', 'destinatarios', fallback='').split(',') if d.strip()],
    }
    if not (config['servidor'] and config['usuario'] and config['senha'] and config['destinatarios']):
        raise ValueError(f"preencha servidor, usuario, senha e destinatarios em {caminho}")
    return config


def enviar_email(config, assunto, corpo, anexos=(), html=None):
    msg = EmailMessage()
    msg['From'] = config['usuario']
    msg['To'] = ', '.join(config['destinatarios'])
    msg['Subject'] = assunto
    msg.set_content(corpo)
    if html:
        msg.add_alternative(html, subtype='html')   # o texto puro fica para quem não abre HTML
    for caminho in anexos:
        with open(caminho, 'rb') as fp:
            msg.add_attachment(fp.read(), maintype='application',
                               subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                               filename=os.path.basename(caminho))
    if config['porta'] == 465:
        conexao = smtplib.SMTP_SSL(config['servidor'], config['porta'], timeout=60)
    else:
        conexao = smtplib.SMTP(config['servidor'], config['porta'], timeout=60)
        conexao.starttls()
    with conexao:
        try:
            conexao.login(config['usuario'], config['senha'])
        except (smtplib.SMTPAuthenticationError, smtplib.SMTPServerDisconnected) as e:
            # o Gmail recusa (ou simplesmente derruba a conexão) quando recebe a
            # senha normal da conta em vez de uma "senha de app"
            raise ValueError(f"login recusado pelo servidor de e-mail ({e}). No Gmail/Google Workspace, "
                             "use uma senha de app de 16 letras (https://myaccount.google.com/apppasswords), "
                             "não a senha normal da conta") from e
        conexao.send_message(msg)


def testar_avisos(pasta_script):
    """Manda uma notificação e um e-mail de teste, para conferir a configuração."""
    titulo = "Pautas CARF/DRJ: teste de aviso"
    corpo = ("Este é um e-mail de teste do script de pautas CARF/DRJ.\n\n"
             "Se você recebeu, os avisos do modo automático estão funcionando.")
    codigo = 0
    try:
        config = ler_config_email(pasta_script)
        enviar_email(config, titulo, corpo)
        print(f"[OK] E-mail de teste enviado para {', '.join(config['destinatarios'])}.")
    except Exception as e:
        print(f"[ERRO] E-mail não enviado: {e}")
        codigo = 2
    notificar_mac(titulo, "Se você está vendo isto, as notificações funcionam.")
    return codigo


def _caminho_plist():
    return os.path.expanduser(f'~/Library/LaunchAgents/{ROTULO_LAUNCHD}.plist')


def instalar_agendamento(pasta_script):
    os.makedirs(os.path.join(pasta_script, PASTA_LOGS), exist_ok=True)
    log = os.path.join(pasta_script, PASTA_LOGS, 'automatico.log')
    config = {
        'Label': ROTULO_LAUNCHD,
        'ProgramArguments': [sys.executable, '-u', os.path.abspath(__file__), '--automatico'],
        'WorkingDirectory': pasta_script,
        'StartCalendarInterval': [{'Weekday': d, 'Hour': h, 'Minute': m}
                                  for d in DIAS_UTEIS for h, m in HORARIOS],
        'StandardOutPath': log,
        'StandardErrorPath': log,
    }
    caminho = _caminho_plist()
    os.makedirs(os.path.dirname(caminho), exist_ok=True)
    dominio = f'gui/{os.getuid()}'
    subprocess.run(['launchctl', 'bootout', f'{dominio}/{ROTULO_LAUNCHD}'], capture_output=True)
    with open(caminho, 'wb') as fp:
        plistlib.dump(config, fp)
    r = subprocess.run(['launchctl', 'bootstrap', dominio, caminho], capture_output=True, text=True)
    if r.returncode != 0:
        print(f"[ERRO] Não consegui ativar o agendamento: {r.stderr.strip() or r.returncode}")
        return 2
    horarios = [f'{h}h' + (f'{m:02d}' if m else '') for h, m in HORARIOS]
    horarios = ', '.join(horarios[:-1]) + ' e ' + horarios[-1] if len(horarios) > 1 else horarios[0]
    print(f"[OK] Agendamento instalado: de segunda a sexta, às {horarios}.")
    print(f"     Log de cada execução: {log}")
    print(f"     Para rodar agora, sem esperar o horário:  launchctl kickstart {dominio}/{ROTULO_LAUNCHD}")
    print(f"     Para desligar:  python3 \"{os.path.basename(__file__)}\" --remover-agendamento")
    return 0


def remover_agendamento():
    subprocess.run(['launchctl', 'bootout', f'gui/{os.getuid()}/{ROTULO_LAUNCHD}'], capture_output=True)
    if os.path.exists(_caminho_plist()):
        os.remove(_caminho_plist())
    print("[OK] Agendamento removido — o script não roda mais sozinho.")
    return 0


# ─────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='Extrai pautas CARF/DRJ de XMLs do INLabs e gera Excel.')
    parser.add_argument('entrada', nargs='?', help='Pasta com XMLs/ZIPs ou caminho direto para um XML')
    parser.add_argument('saida',   nargs='?', help='Arquivo Excel de saída (.xlsx)')
    parser.add_argument('--clientes', dest='clientes',
                         help='Planilha (.xlsx/.csv) com os clientes do escritório, para marcar '
                              'quando algum aparecer na pauta (colunas Nome e, opcionalmente, CNPJ/CPF)')
    parser.add_argument('--baixar', nargs='?', const=date.today().isoformat(), type=interpretar_data,
                        metavar='DATA',
                        help='Baixa do INLabs todas as seções do DOU da data (padrão: hoje; '
                             'AAAA-MM-DD ou DD/MM/AAAA) e processa. A planilha sai em '
                             'pautas_<data>.xlsx, na pasta do script; os ZIPs baixados são apagados no final.')
    parser.add_argument('--manter-dou', action='store_true',
                        help='Com --baixar: guarda os ZIPs baixados em DOU/<data>/ em vez de apagar.')
    grupo = parser.add_argument_group('modo automático')
    grupo.add_argument('--automatico', action='store_true',
                       help='Processa hoje e os dias pendentes e avisa (notificação + e-mail) se aparecer '
                            'cliente ou der problema. É o que o agendamento roda.')
    grupo.add_argument('--instalar-agendamento', action='store_true',
                       help='Faz o Mac rodar o modo automático sozinho, de segunda a sexta às 8h30, 13h e 18h.')
    grupo.add_argument('--remover-agendamento', action='store_true', help='Desliga o agendamento.')
    grupo.add_argument('--testar-aviso', action='store_true',
                       help='Manda uma notificação e um e-mail de teste.')
    args = parser.parse_args()

    pasta_script = os.path.dirname(os.path.abspath(__file__))
    if args.automatico:
        sys.exit(rodar_automatico(pasta_script))
    if args.instalar_agendamento:
        sys.exit(instalar_agendamento(pasta_script))
    if args.remover_agendamento:
        sys.exit(remover_agendamento())
    if args.testar_aviso:
        sys.exit(testar_avisos(pasta_script))

    if args.manter_dou and not args.baixar:
        parser.error('--manter-dou só faz sentido junto com --baixar')
    if args.baixar and args.entrada and args.saida:
        parser.error('com --baixar, informe no máximo o arquivo de saída (a entrada é o próprio download)')

    if not (args.entrada or args.saida or args.baixar):
        # rodado sem opções (ex.: botão "Run" do VS Code): pergunta o que fazer
        args.baixar = menu_inicial()

    sys.exit(executar(args, pasta_script).codigo)


if __name__ == '__main__':
    main()
