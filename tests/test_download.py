# -*- coding: utf-8 -*-
"""
Testes do download do INLabs contra um servidor falso local (não usa a
internet nem a senha real).

Rodar (da pasta do projeto):
    python3 -m unittest discover tests -v
"""

import contextlib
import http.server
import importlib.util
import io
import os
import tempfile
import threading
import unittest
import urllib.parse
import zipfile
from unittest import mock

from openpyxl import Workbook, load_workbook

PASTA_TESTES = os.path.dirname(os.path.abspath(__file__))
PASTA_FIXTURES = os.path.join(PASTA_TESTES, 'fixtures')
SCRIPT = os.path.join(PASTA_TESTES, '..', 'pautas_carf.py')

_spec = importlib.util.spec_from_file_location('pauta_download', SCRIPT)
pauta = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pauta)

DATA = '2026-09-25'
EMAIL, SENHA = 'teste@exemplo.com', 'senha%com%porcento'


def _zip_com(*nomes_fixture):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as zf:
        for nome in nomes_fixture:
            zf.write(os.path.join(PASTA_FIXTURES, nome), nome)
    return buf.getvalue()


ZIP_DO1 = _zip_com('515_20260925_24408020.xml', 'nao_pauta.xml')   # 3 processos
ZIP_VAZIO_DE_PAUTA = _zip_com('nao_pauta.xml')


class InlabsFalso(http.server.BaseHTTPRequestHandler):
    """
    Comportamento por seção em `comportamento`: 'ok', 'ok_sem_pauta', '404',
    '500' (sempre), 'html' (200 sem ser ZIP), 'login' (página de login: sessão
    expirada), 'listagem' (página logada sem a data: é o que o INLabs devolve
    para fim de semana/feriado) ou 'instavel' (500 na 1ª vez).
    """
    comportamento = {}
    chamadas = []
    logins = 0
    logins_falhando = 0   # quantos logins seguidos o INLabs recusa mesmo com a senha certa

    def log_message(self, *args):
        pass

    def do_POST(self):
        corpo = self.rfile.read(int(self.headers['Content-Length']))
        dados = urllib.parse.parse_qs(corpo.decode())
        InlabsFalso.logins += 1
        self.send_response(302)
        if (dados.get('email') == [EMAIL] and dados.get('password') == [SENHA]
                and InlabsFalso.logins > InlabsFalso.logins_falhando):
            self.send_header('Set-Cookie', 'inlabs_session_cookie=abc123; Path=/')
        self.send_header('Location', '/index.php')
        self.end_headers()

    def do_GET(self):
        if not self.path.startswith('/index.php?'):
            return self._responder(200, b'<html>inicio</html>')
        if 'inlabs_session_cookie=abc123' not in (self.headers.get('Cookie') or ''):
            return self._responder(200, b'<html>faca login</html>')
        dl = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)['dl'][0]
        secao = dl.rsplit('-', 1)[1].replace('.zip', '')
        InlabsFalso.chamadas.append(secao)
        modo = self.comportamento.get(secao, '404')
        if modo == 'instavel':
            modo = '500' if InlabsFalso.chamadas.count(secao) == 1 else 'ok'
        if modo == 'ok':
            self._responder(200, ZIP_DO1)
        elif modo == 'ok_sem_pauta':
            self._responder(200, ZIP_VAZIO_DE_PAUTA)
        elif modo == 'html':
            self._responder(200, b'<html>sessao expirada</html>')
        elif modo == 'login':
            self._responder(200, b'<!DOCTYPE html><html><form>E-mail: <input name="email"> '
                                 b'Senha: <input type="password" name="password"></form></html>')
        elif modo == 'listagem':
            self._responder(200, b'<!DOCTYPE html><html>Ola. Minha Conta Ajuda Sair '
                                 b'<a>2026-09-01</a> <a>2026-08-31</a></html>')
        else:
            self._responder(int(modo), b'')

    def _responder(self, status, corpo):
        self.send_response(status)
        self.send_header('Content-Length', str(len(corpo)))
        self.end_headers()
        self.wfile.write(corpo)


class BaseDownload(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.servidor = http.server.ThreadingHTTPServer(('127.0.0.1', 0), InlabsFalso)
        threading.Thread(target=cls.servidor.serve_forever, daemon=True).start()
        base = f'http://127.0.0.1:{cls.servidor.server_port}'
        cls.patches = [
            mock.patch.object(pauta, 'URL_LOGIN', base + '/logar.php'),
            mock.patch.object(pauta, 'URL_DOWNLOAD', base + '/index.php?p={data}&dl={data}-{secao}.zip'),
            mock.patch.object(pauta, 'ESPERA_ENTRE_TENTATIVAS', 0),
        ]
        for p in cls.patches:
            p.start()

    @classmethod
    def tearDownClass(cls):
        for p in cls.patches:
            p.stop()
        cls.servidor.shutdown()
        cls.servidor.server_close()

    def setUp(self):
        # garante que nenhum teste fala com o INLabs de verdade
        self.assertTrue(pauta.URL_LOGIN.startswith('http://127.0.0.1:'), 'teste apontando para o INLabs real!')
        InlabsFalso.chamadas = []
        InlabsFalso.comportamento = {}
        InlabsFalso.logins = InlabsFalso.logins_falhando = 0
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.env = mock.patch.dict(os.environ, {'INLABS_EMAIL': EMAIL, 'INLABS_SENHA': SENHA})
        self.env.start()
        self.addCleanup(self.env.stop)

    def baixar(self, **comportamento):
        InlabsFalso.comportamento = comportamento
        with contextlib.redirect_stdout(io.StringIO()):
            return pauta.baixar_dou(DATA, os.path.join(self.tmp.name, 'DOU'), self.tmp.name)

    def situacoes(self, resultados):
        return {secao: situacao for secao, situacao, _ in resultados}


class TestBaixarDou(BaseDownload):

    def test_baixa_todas_as_secoes(self):
        todas = {s: 'ok' for s in pauta.SECOES_DOU}
        resultados = self.baixar(**todas)
        self.assertEqual(self.situacoes(resultados), {s: 'OK' for s in pauta.SECOES_DOU})
        self.assertEqual(InlabsFalso.chamadas, pauta.SECOES_DOU)
        for secao in pauta.SECOES_DOU:
            self.assertTrue(zipfile.is_zipfile(os.path.join(self.tmp.name, 'DOU', f'{DATA}-{secao}.zip')))

    def test_situacoes_de_cada_secao(self):
        resultados = self.baixar(DO1='ok', DO1E='404', DO2='404', DO2E='500', DO3='html', DO3E='ok')
        self.assertEqual(self.situacoes(resultados), {
            'DO1': 'OK', 'DO1E': 'SEM_EXTRA', 'DO2': 'NAO_PUBLICADO',
            'DO2E': 'ERRO', 'DO3': 'ERRO', 'DO3E': 'OK',
        })
        detalhes = {secao: detalhe for secao, _, detalhe in resultados}
        self.assertIn('3 tentativas', detalhes['DO2E'])
        self.assertIn('HTTP 500', detalhes['DO2E'])
        self.assertIn('ZIP', detalhes['DO3'])
        # nada de arquivo quebrado ou .part largado para trás
        self.assertEqual(sorted(os.listdir(os.path.join(self.tmp.name, 'DOU'))),
                         [f'{DATA}-DO1.zip', f'{DATA}-DO3E.zip'])

    def test_data_sem_dou_vem_como_pagina_de_listagem(self):
        # fim de semana/feriado: o INLabs devolve a lista de datas (logado), não 404
        resultados = self.baixar(DO1='listagem', DO1E='listagem', DO2='login', DO3='ok')
        self.assertEqual(self.situacoes(resultados)['DO1'], 'NAO_PUBLICADO')
        self.assertEqual(self.situacoes(resultados)['DO1E'], 'SEM_EXTRA')
        self.assertEqual(self.situacoes(resultados)['DO2'], 'ERRO')      # página de login: sessão expirada
        self.assertEqual(self.situacoes(resultados)['DO3'], 'OK')

    def test_listagem_que_tem_a_data_continua_sendo_erro(self):
        self.assertFalse(pauta._pagina_sem_a_data(b'<!DOCTYPE html><html>Sair 2026-10-02</html>', '2026-10-02'))
        self.assertTrue(pauta._pagina_sem_a_data(b'<!DOCTYPE html><html>Sair 2026-10-02</html>', '2026-10-03'))
        self.assertFalse(pauta._pagina_sem_a_data(b'PK\x03\x04', '2026-10-03'))

    def test_falha_temporaria_e_repetida(self):
        resultados = self.baixar(DO1='instavel')
        self.assertEqual(self.situacoes(resultados)['DO1'], 'OK')
        self.assertEqual(InlabsFalso.chamadas.count('DO1'), 2)

    def test_nao_baixa_de_novo_o_que_ja_esta_baixado(self):
        self.baixar(DO1='ok')
        InlabsFalso.chamadas = []
        resultados = self.baixar(DO1='ok')
        self.assertNotIn('DO1', InlabsFalso.chamadas)
        self.assertIn('já estava baixado', dict((s, d) for s, _, d in resultados)['DO1'])

    def test_edicao_extra_sempre_baixa_de_novo(self):
        # uma edição extra pode sair mais tarde no mesmo dia
        self.baixar(DO1='ok', DO1E='ok')
        InlabsFalso.chamadas = []
        self.baixar(DO1='ok', DO1E='ok')
        self.assertIn('DO1E', InlabsFalso.chamadas)
        self.assertNotIn('DO1', InlabsFalso.chamadas)

    def test_tamanho_legivel(self):
        caminho = os.path.join(self.tmp.name, 'x')
        for n, esperado in ((1946, '2 KB'), (9023, '9 KB'), (2_560_654, '2,6 MB'), (10, '1 KB')):
            with open(caminho, 'wb') as fp:
                fp.write(b'0' * n)
            self.assertEqual(pauta._tamanho(caminho), esperado)

    def test_senha_errada(self):
        with mock.patch.dict(os.environ, {'INLABS_SENHA': 'errada'}):
            with self.assertRaisesRegex(pauta.ErroInlabs, 'recusou o login'):
                self.baixar(DO1='ok')
        self.assertEqual(InlabsFalso.chamadas, [])
        self.assertEqual(InlabsFalso.logins, 5)

    def test_login_instavel_entra_numa_tentativa_seguinte(self):
        InlabsFalso.logins_falhando = 4
        resultados = self.baixar(DO1='ok')
        self.assertEqual(InlabsFalso.logins, 5)
        self.assertEqual(self.situacoes(resultados)['DO1'], 'OK')

    def test_servidor_fora_do_ar(self):
        with mock.patch.object(pauta, 'URL_LOGIN', 'http://127.0.0.1:9/logar.php'):
            with self.assertRaisesRegex(pauta.ErroInlabs, 'Não consegui acessar o INLabs'):
                self.baixar(DO1='ok')


class TestCredenciais(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        env = mock.patch.dict(os.environ, {}, clear=False)
        env.start()
        self.addCleanup(env.stop)
        os.environ.pop('INLABS_EMAIL', None)
        os.environ.pop('INLABS_SENHA', None)
        self.caminho = os.path.join(self.tmp.name, pauta.ARQ_CREDENCIAIS)

    def test_cria_modelo_na_primeira_vez(self):
        with self.assertRaisesRegex(pauta.ErroInlabs, 'Criei o arquivo'):
            pauta.ler_credenciais(self.tmp.name)
        self.assertTrue(os.path.exists(self.caminho))
        self.assertEqual(os.stat(self.caminho).st_mode & 0o777, 0o600)  # só o dono lê

    def test_modelo_nao_preenchido(self):
        with open(self.caminho, 'w') as fp:
            fp.write(pauta.MODELO_CREDENCIAIS)
        with self.assertRaisesRegex(pauta.ErroInlabs, 'Preencha'):
            pauta.ler_credenciais(self.tmp.name)

    def test_le_arquivo_com_senha_com_porcento(self):
        with open(self.caminho, 'w') as fp:
            fp.write(f'[inlabs]\nemail = {EMAIL}\nsenha = {SENHA}\n')
        self.assertEqual(pauta.ler_credenciais(self.tmp.name), (EMAIL, SENHA))


class TestMenuInicial(unittest.TestCase):

    def menu(self, *respostas):
        with mock.patch('builtins.input', side_effect=list(respostas)), \
             contextlib.redirect_stdout(io.StringIO()) as saida:
            return pauta.menu_inicial(), saida.getvalue()

    def test_enter_baixa_hoje(self):
        self.assertEqual(self.menu('')[0], pauta.date.today().isoformat())

    def test_outra_data(self):
        self.assertEqual(self.menu('25/09/2026')[0], '2026-09-25')

    def test_escolher_pasta(self):
        self.assertIsNone(self.menu('P')[0])

    def test_resposta_invalida_pergunta_de_novo(self):
        resultado, saida = self.menu('ontem', '31/02/2026', '2026-09-25')
        self.assertEqual(resultado, '2026-09-25')
        self.assertEqual(saida.count('tente de novo'), 2)

    def test_sem_terminal_encerra(self):
        with mock.patch('builtins.input', side_effect=EOFError), \
             contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(SystemExit) as cm:
                pauta.menu_inicial()
        self.assertEqual(cm.exception.code, 0)


class TestInterpretarData(unittest.TestCase):

    def test_formatos(self):
        self.assertEqual(pauta.interpretar_data('2026-09-25'), '2026-09-25')
        self.assertEqual(pauta.interpretar_data('25/09/2026'), '2026-09-25')
        with self.assertRaises(Exception):
            pauta.interpretar_data('31/02/2026')


class TestMainComDownload(BaseDownload):
    """Roda o main() de verdade com --baixar, numa pasta temporária."""

    LISTA_CLIENTES = 'Clientes monitorados.xlsx'

    def setUp(self):
        super().setUp()
        wb = Workbook()
        wb.active.append(['Nome', 'Processo'])
        wb.active.append(['Cliente Ficticio Ltda', '11111.111111/2020-11'])
        wb.save(os.path.join(self.tmp.name, self.LISTA_CLIENTES))

    def rodar_main(self, *args_extras, **comportamento):
        InlabsFalso.comportamento = comportamento
        saida = io.StringIO()
        with mock.patch.object(pauta, '__file__', os.path.join(self.tmp.name, 'script.py')), \
             mock.patch('sys.argv', ['script.py', '--baixar', '25/09/2026', *args_extras]), \
             contextlib.redirect_stdout(saida):
            try:
                pauta.main()
                codigo = 0
            except SystemExit as e:
                codigo = e.code
        return codigo, saida.getvalue()

    def test_tudo_certo(self):
        codigo, saida = self.rodar_main(DO1='ok', DO2='ok_sem_pauta', DO3='ok_sem_pauta')
        self.assertEqual(codigo, 0, saida)
        self.assertNotIn('ATENÇÃO', saida)
        self.assertIn('[--]    DO1E  sem edição extra', saida)
        self.assertTrue(os.path.exists(os.path.join(self.tmp.name, f'pautas_{DATA}.xlsx')))

    def test_erro_em_uma_secao_avisa_no_final_e_sai_com_codigo_2(self):
        codigo, saida = self.rodar_main(DO1='ok', DO2='500', DO3='ok_sem_pauta')
        self.assertEqual(codigo, 2)
        self.assertIn('[ERRO]  DO2', saida)
        # a planilha é gerada com o que deu para baixar...
        self.assertTrue(os.path.exists(os.path.join(self.tmp.name, f'pautas_{DATA}.xlsx')))
        # ...e o aviso aparece por último, depois do "[OK] Planilha salva"
        final = saida[saida.index('[OK] Planilha salva'):]
        self.assertIn('[ATENÇÃO] A planilha pode estar incompleta', final)
        self.assertIn('1 seção(ões) do DOU NÃO foram baixadas', final)
        self.assertNotIn('Clientes NÃO foram checados', final)
        self.assertIn('- DO2:', final)
        self.assertIn('--baixar 2026-09-25', final)

    def test_senha_errada_para_com_codigo_2(self):
        with mock.patch.dict(os.environ, {'INLABS_SENHA': 'errada'}):
            codigo, saida = self.rodar_main(DO1='ok')
        self.assertEqual(codigo, 2)
        self.assertIn('[ERRO] O INLabs recusou o login', saida)

    def pasta_dou(self):
        return os.path.join(self.tmp.name, 'DOU')

    def test_apaga_os_zips_e_fica_so_a_planilha(self):
        codigo, saida = self.rodar_main(DO1='ok', DO2='ok_sem_pauta')
        self.assertEqual(codigo, 0, saida)
        self.assertFalse(os.path.exists(self.pasta_dou()))
        self.assertEqual(sorted(os.listdir(self.tmp.name)), [self.LISTA_CLIENTES, f'pautas_{DATA}.xlsx'])
        self.assertIn('ZIPs do DOU apagados', saida)

    def test_manter_dou_guarda_os_zips(self):
        codigo, saida = self.rodar_main('--manter-dou', DO1='ok', DO2='ok_sem_pauta')
        self.assertEqual(codigo, 0, saida)
        self.assertEqual(sorted(os.listdir(os.path.join(self.pasta_dou(), DATA))),
                         [f'{DATA}-DO1.zip', f'{DATA}-DO2.zip'])
        self.assertIn('ZIPs do DOU mantidos em', saida)

    def test_apaga_tambem_quando_da_erro_no_download(self):
        codigo, saida = self.rodar_main(DO1='ok', DO2='500')
        self.assertEqual(codigo, 2)
        self.assertFalse(os.path.exists(self.pasta_dou()))
        # o aviso de erro continua sendo a última coisa da saída
        self.assertTrue(saida.rstrip().endswith('!' * 70))

    def test_senha_errada_nao_deixa_pasta_vazia(self):
        with mock.patch.dict(os.environ, {'INLABS_SENHA': 'errada'}):
            self.rodar_main(DO1='ok')
        self.assertFalse(os.path.exists(self.pasta_dou()))

    def test_erro_inesperado_mantem_os_zips(self):
        with mock.patch.object(pauta, 'gerar_xlsx', side_effect=RuntimeError('disco cheio')):
            with self.assertRaises(RuntimeError):
                self.rodar_main(DO1='ok')
        self.assertTrue(os.path.exists(os.path.join(self.pasta_dou(), DATA, f'{DATA}-DO1.zip')))

    def test_sem_lista_de_clientes_avisa_no_final(self):
        os.remove(os.path.join(self.tmp.name, self.LISTA_CLIENTES))
        codigo, saida = self.rodar_main(DO1='ok')
        self.assertEqual(codigo, 2)
        final = saida[saida.index('[OK] Planilha salva'):]
        self.assertIn('Clientes NÃO foram checados: nenhuma lista de clientes', final)

    def test_lista_de_clientes_ilegivel_avisa_no_final(self):
        with open(os.path.join(self.tmp.name, self.LISTA_CLIENTES), 'w') as fp:
            fp.write('isto não é um xlsx')
        codigo, saida = self.rodar_main(DO1='ok')
        self.assertEqual(codigo, 2)
        self.assertIn('Clientes NÃO foram checados: não consegui ler a lista de clientes', saida)

    def test_dia_sem_pautas_gera_planilha_vazia_sem_erro(self):
        codigo, saida = self.rodar_main(DO1='ok_sem_pauta', DO2='ok_sem_pauta', DO3='ok_sem_pauta')
        self.assertEqual(codigo, 0, saida)
        self.assertIn('Nenhuma pauta de julgamento do CARF/DRJ', saida)
        self.assertNotIn('ATENÇÃO', saida)
        wb = load_workbook(os.path.join(self.tmp.name, f'pautas_{DATA}.xlsx'))
        self.assertEqual(wb['Pautas'].max_row, 1)   # só o cabeçalho
        self.assertEqual(wb['Pautas'].cell(row=1, column=1).value, 'Órgão')

    def test_dia_sem_dou(self):
        codigo, saida = self.rodar_main()  # tudo 404
        self.assertEqual(codigo, 1)
        self.assertIn('[AVISO] DO1   não publicado', saida)
        self.assertIn('nada a processar', saida)
        self.assertFalse(os.path.exists(self.pasta_dou()))


if __name__ == '__main__':
    unittest.main()
